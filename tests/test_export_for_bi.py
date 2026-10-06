import re

import numpy as np
import pandas as pd
import pytest
from openpyxl import load_workbook

import export_for_bi as bi
from forecaster import config


def make_predictions(hours=24 * 365, start="2025-01-01", with_weather=True):
    """Fake predictions with the same columns as outputs/predictions.csv."""
    idx = pd.date_range(start, periods=hours, freq="h", name="timestamp")
    rng = np.random.default_rng(0)
    actual = (15000 + 3000 * np.sin(2 * np.pi * idx.hour / 24) + rng.normal(0, 300, hours)).round()
    df = pd.DataFrame({"ontario_demand": actual,
                       "pred_xgboost": actual + rng.normal(0, 400, hours),
                       "pred_naive_yesterday": actual + rng.normal(0, 700, hours),
                       "pred_naive_last_week": actual + rng.normal(0, 1100, hours)}, index=idx)
    if with_weather:
        df["temperature_c"] = (8 - 14 * np.cos(2 * np.pi * idx.dayofyear / 365)).round(1)
    return df


@pytest.fixture(scope="module")
def exported(tmp_path_factory):
    """Export a full fake year once and read everything back from the files."""
    out = tmp_path_factory.mktemp("bi")
    preds = make_predictions()
    bi.export(preds, out)
    csv = pd.read_csv(out / bi.CSV_NAME, parse_dates=["timestamp"])
    formulas = load_workbook(out / bi.XLSX_NAME)                    # cells hold formula text
    cached = load_workbook(out / bi.XLSX_NAME, data_only=True)      # cells hold cached results
    rows = list(formulas["Data"].iter_rows(values_only=True))
    data = pd.DataFrame(rows[1:], columns=rows[0])
    return {"preds": preds, "csv": csv, "formulas": formulas, "cached": cached, "data": data}


def evaluate(formula, ws, data):
    """Work out a COUNT / AVERAGE / COUNTIFS / AVERAGEIFS formula by hand from the
    Data sheet's rows, so the formulas are checked without needing Excel."""
    body = formula.lstrip("=")
    if body.startswith("IFERROR("):
        body = body[len("IFERROR("):body.rindex(",")]
    name, args = body.rstrip(")").split("(", 1)
    args = [a[len(bi.TABLE) + 1:-1] if a.startswith(bi.TABLE + "[")
            else ws[a.replace("$", "")].value for a in args.split(",")]
    values = None if name == "COUNTIFS" else data[args.pop(0)]
    mask = pd.Series(True, index=data.index)
    for column, wanted in zip(args[::2], args[1::2]):
        mask &= data[column] == wanted
    if name in ("COUNT", "COUNTIFS"):
        return int(mask.sum())
    assert name in ("AVERAGE", "AVERAGEIFS"), f"unexpected formula: {formula}"
    return values[mask].mean()


# ---- the table --------------------------------------------------------------
def test_errors_are_computed_from_actual_predicted_and_baseline():
    idx = pd.date_range("2025-03-03 14:00", periods=1, freq="h")       # a Monday
    preds = pd.DataFrame({"ontario_demand": [100.0], "pred_xgboost": [110.0],
                          "pred_naive_yesterday": [80.0], "temperature_c": [25.0]}, index=idx)
    row = bi.build_results(preds).iloc[0]
    assert (row["abs_error_mw"], row["baseline_abs_error_mw"]) == (10.0, 20.0)
    assert (row["abs_pct_error"], row["baseline_abs_pct_error"]) == (0.1, 0.2)
    assert (row["hour"], row["month"], row["weekday"]) == (14, 3, "Monday")
    assert (row["heating_degrees"], row["cooling_degrees"]) == (0.0, 7.0)


def test_one_row_per_prediction(exported):
    n = len(exported["preds"])
    assert len(exported["csv"]) == n
    assert len(exported["data"]) == n
    assert exported["csv"]["timestamp"].is_unique


def test_no_nulls_in_required_columns(exported):
    for table in (exported["csv"], exported["data"]):
        assert table[["timestamp", *bi.REQUIRED]].notna().all().all()


def test_csv_and_workbook_hold_the_same_numbers(exported):
    csv, data = exported["csv"], exported["data"]
    assert list(csv.columns) == list(data.columns)
    numeric = csv.select_dtypes("number").columns
    assert np.allclose(csv[numeric], data[numeric].astype(float), rtol=0, atol=1e-9)
    assert (csv["weekday"] == data["weekday"]).all()
    assert (csv["timestamp"] == pd.to_datetime(data["timestamp"]).dt.round("s")).all()


def test_export_works_without_weather(tmp_path):
    results = bi.export(make_predictions(hours=48, with_weather=False), tmp_path)
    assert "temperature_c" not in results.columns


def test_export_rejects_missing_values(tmp_path):
    preds = make_predictions(hours=48)
    preds.iloc[5, preds.columns.get_loc("pred_xgboost")] = np.nan
    with pytest.raises(ValueError, match="predicted_mw"):
        bi.export(preds, tmp_path)


# ---- the workbook -----------------------------------------------------------
def test_data_sheet_is_an_excel_table(exported):
    ws = exported["formulas"]["Data"]
    n_rows, n_cols = len(exported["data"]) + 1, len(exported["data"].columns)
    assert ws.tables[bi.TABLE].ref == f"A1:{ws.cell(n_rows, n_cols).coordinate}"


def test_overall_summary_matches_pandas(exported):
    ws, cached, csv = exported["formulas"]["Summary"], exported["cached"]["Summary"], exported["csv"]
    model, base = csv["abs_pct_error"].mean(), csv["baseline_abs_pct_error"].mean()
    model_mw, base_mw = csv["abs_error_mw"].mean(), csv["baseline_abs_error_mw"].mean()
    by_hour = csv.groupby("hour")[["abs_pct_error", "abs_error_mw"]].mean()
    by_month = csv.groupby("month")["abs_pct_error"].mean()

    # Formulas over the table, worked out by hand.
    for cell, expected in [("D5", len(csv)), ("D6", model), ("D7", base),
                           ("D9", model_mw), ("D10", base_mw)]:
        assert evaluate(ws[cell].value, ws, exported["data"]) == pytest.approx(expected)
        assert cached[cell].value == pytest.approx(expected)

    # Formulas that build on other cells: check the formula text and the cached result.
    assert ws["D8"].value == "=(D7-D6)/D7"
    assert ws["D11"].value == "=(D10-D9)/D10"
    assert cached["D8"].value == pytest.approx((base - model) / base)
    assert cached["D11"].value == pytest.approx((base_mw - model_mw) / base_mw)
    for cell in ("D12", "D13", "D14", "D15"):
        assert ws[cell].value.startswith("=INDEX(")
    assert cached["D12"].value == by_hour["abs_pct_error"].idxmax()
    assert cached["D13"].value == by_hour["abs_error_mw"].idxmax()
    assert cached["D14"].value == pd.Timestamp(2025, by_month.idxmax(), 1).strftime("%b")
    assert cached["D15"].value == pd.Timestamp(2025, by_month.idxmin(), 1).strftime("%b")


@pytest.mark.parametrize("column, first_row, keys", [("month", bi.MONTH_ROW, range(1, 13)),
                                                     ("hour", bi.HOUR_ROW, range(24))])
def test_grouped_summary_matches_pandas(exported, column, first_row, keys):
    ws, cached, csv = exported["formulas"]["Summary"], exported["cached"]["Summary"], exported["csv"]
    g = csv.groupby(column)
    expected = pd.DataFrame({"C": g.size(), "D": g["abs_pct_error"].mean(),
                             "E": g["baseline_abs_pct_error"].mean(),
                             "G": g["abs_error_mw"].mean()})
    for i, key in enumerate(keys):
        r = first_row + i
        assert ws[f"A{r}"].value == key
        for col in "CDEG":
            formula = ws[f"{col}{r}"].value
            assert ("COUNTIFS(" if col == "C" else "AVERAGEIFS(") in formula
            assert evaluate(formula, ws, exported["data"]) == pytest.approx(expected.loc[key, col])
            assert cached[f"{col}{r}"].value == pytest.approx(expected.loc[key, col])
        e, d = expected.loc[key, "E"], expected.loc[key, "D"]
        assert ws[f"F{r}"].value == f'=IFERROR((E{r}-D{r})/E{r},"")'
        assert cached[f"F{r}"].value == pytest.approx((e - d) / e)


def test_heatmap_is_formulas_with_a_colour_scale(exported):
    ws, cached, csv = exported["formulas"]["Heatmap"], exported["cached"]["Heatmap"], exported["csv"]
    expected = csv.pivot_table(index="hour", columns="month", values="abs_pct_error")
    for hour in range(24):
        for month in range(1, 13):
            cell = ws.cell(bi.HEAT_ROW + hour, month + 1)
            assert "AVERAGEIFS(" in cell.value
            assert evaluate(cell.value, ws, exported["data"]) == pytest.approx(expected.loc[hour, month])
            assert cached[cell.coordinate].value == pytest.approx(expected.loc[hour, month])
    scales = [rule.type for cf in ws.conditional_formatting for rule in cf.rules]
    assert scales == ["colorScale"]


def test_peak_week_chart_plots_actual_and_predicted_for_one_week(exported):
    chart, = exported["formulas"]["Peak week"]._charts
    refs = [s.val.numRef.f for s in chart.series]
    assert len(refs) == 2
    data = exported["data"]
    for ref, col in zip(refs, ("actual_mw", "predicted_mw")):
        letter, first, last = re.fullmatch(r"'?Data'?!\$([A-Z]+)\$(\d+):\$[A-Z]+\$(\d+)", ref).groups()
        assert exported["formulas"]["Data"][f"{letter}1"].value == col
        assert int(last) - int(first) + 1 == 24 * 7
        # The plotted week contains the highest-demand hour.
        assert int(first) - 2 <= data["actual_mw"].idxmax() <= int(last) - 2


def test_months_without_data_stay_blank(tmp_path):
    preds = make_predictions(hours=24 * 10, start="2025-06-10")
    results = bi.export(preds, tmp_path)
    stats = bi.summarize(results)
    assert stats["by_month"].loc[1, "hours"] == 0
    assert stats["hardest_month"] == 6
    start, end = bi.peak_week(results)
    assert results.index.min() <= start and end <= results.index.max()

    cached = load_workbook(tmp_path / bi.XLSX_NAME, data_only=True)["Summary"]
    assert cached[f"C{bi.MONTH_ROW}"].value == 0                         # January: no hours
    assert cached[f"D{bi.MONTH_ROW}"].value in (None, "")                # blank, not #DIV/0!
    assert cached[f"D{bi.MONTH_ROW + 5}"].value == pytest.approx(results["abs_pct_error"].mean())


# ---- the script -------------------------------------------------------------
def test_main_reads_outputs_and_writes_bi_files(tmp_path, monkeypatch, capsys):
    monkeypatch.setattr(config, "OUTPUTS", tmp_path / "outputs")
    monkeypatch.setattr(bi, "BI_DIR", tmp_path / "bi")
    with pytest.raises(SystemExit, match="forecaster train"):
        bi.main()                                                        # nothing trained yet

    (tmp_path / "outputs").mkdir()
    make_predictions(hours=24 * 14).to_csv(tmp_path / "outputs" / "predictions.csv")
    bi.main()
    assert (tmp_path / "bi" / bi.CSV_NAME).exists()
    assert (tmp_path / "bi" / bi.XLSX_NAME).exists()
    assert "336 rows" in capsys.readouterr().out
